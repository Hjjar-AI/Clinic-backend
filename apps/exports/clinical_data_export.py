# backend/apps/exports/clinical_data_export.py
from openpyxl import Workbook
from django.http import HttpResponse
from django.db.models import Q
from django.utils import timezone

from apps.clinical.models import DiagnosisOption, MedicationOption
from core.query_utils import apply_ordering

class ClinicalDataExportService:
    @staticmethod
    def _metadata_sheet(workbook, filters):
        sheet = workbook.create_sheet('معلومات التصدير')
        sheet.append(['وقت الإنشاء', timezone.localtime().isoformat()])
        sheet.append(['المنطقة الزمنية', timezone.get_current_timezone_name()])
        for key in ('search', 'activity', 'sort_by', 'sort_order'):
            sheet.append([key, str(filters.get(key, '') or '')])

    def export_diagnoses_excel(self, filters=None, include_inactive=False):
        filters = filters or {}
        diagnoses = DiagnosisOption.all_objects.all() if include_inactive else DiagnosisOption.objects.all()
        search = str(filters.get('search', '')).strip()
        if search:
            diagnoses = diagnoses.filter(
                Q(code__icontains=search)
                | Q(english_name__icontains=search)
                | Q(arabic_name__icontains=search)
            )
        activity = filters.get('activity')
        if activity == 'active':
            diagnoses = diagnoses.filter(is_active=True, deleted_at__isnull=True)
        elif activity == 'inactive' and include_inactive:
            diagnoses = diagnoses.filter(Q(is_active=False) | Q(deleted_at__isnull=False))
        diagnoses = apply_ordering(diagnoses, filters, {
            'order': 'order', 'code': 'code',
            'english_name': 'english_name', 'arabic_name': 'arabic_name',
        }, default='order')
        wb = Workbook()
        ws = wb.active
        ws.title = 'التشخيصات'
        ws.append(['code', 'english_name', 'arabic_name', 'status'])
        for d in diagnoses:
            ws.append([d.code, d.english_name, d.arabic_name, 'active' if d.is_active else 'inactive'])
        self._metadata_sheet(wb, filters)
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="diagnosis_export.xlsx"'
        wb.save(response)
        return response

    def export_medications_excel(self, filters=None, include_inactive=False):
        filters = filters or {}
        meds = MedicationOption.all_objects.all() if include_inactive else MedicationOption.objects.all()
        search = str(filters.get('search', '')).strip()
        if search:
            meds = meds.filter(
                Q(generic_english__icontains=search)
                | Q(generic_arabic__icontains=search)
                | Q(brand_english__icontains=search)
                | Q(brand_arabic__icontains=search)
                | Q(dosage__icontains=search)
            )
        activity = filters.get('activity')
        if activity == 'active':
            meds = meds.filter(is_active=True, deleted_at__isnull=True)
        elif activity == 'inactive' and include_inactive:
            meds = meds.filter(Q(is_active=False) | Q(deleted_at__isnull=False))
        meds = apply_ordering(meds, filters, {
            'order': 'order', 'generic_english': 'generic_english',
            'generic_arabic': 'generic_arabic', 'dosage': 'dosage',
            'brand_english': 'brand_english', 'brand_arabic': 'brand_arabic',
        }, default='order')
        wb = Workbook()
        ws = wb.active
        ws.title = 'الأدوية'
        ws.append(['generic_english', 'generic_arabic', 'dosage', 'brand_english', 'brand_arabic', 'controlled', 'status'])
        for m in meds:
            ws.append([
                m.generic_english, m.generic_arabic or '', m.dosage or '',
                m.brand_english or '', m.brand_arabic or '',
                'yes' if m.is_controlled else 'no', 'active' if m.is_active else 'inactive',
            ])
        self._metadata_sheet(wb, filters)
        response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        response['Content-Disposition'] = 'attachment; filename="medications_export.xlsx"'
        wb.save(response)
        return response
