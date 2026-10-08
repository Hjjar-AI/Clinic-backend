import io
from docx import Document
from docx.shared import Pt, Mm
from docx.enum.text import WD_ALIGN_PARAGRAPH

def export_patient_to_word(patient, clinic_info):
    """
    Generate a Word document for a patient report.
    Returns BytesIO buffer.
    """
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    doc.sections[0].page_width = Mm(210)
    doc.sections[0].page_height = Mm(297)

    # Clinic header
    title = doc.add_heading(clinic_info['name'], 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if clinic_info.get('address') or clinic_info.get('phone'):
        p = doc.add_paragraph(f"{clinic_info.get('address', '')} | {clinic_info.get('phone', '')}")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _add_rtl_paragraph(doc, f'رقم الملف: {patient.patient_number}')
    # Patient information
    doc.add_heading(f"تقرير المريض: {patient.get_full_name()}", level=1)
    doc.add_heading('المعلومات الشخصية', level=2)
    _add_rtl_paragraph(doc, f"الاسم: {patient.get_full_name()}")
    _add_rtl_paragraph(doc, f"سنة الميلاد: {patient.dob_year}")
    _add_rtl_paragraph(doc, f"الجنس: {patient.gender}")
    _add_rtl_paragraph(doc, f"الرقم الوطني: {patient.national_id}")
    _add_rtl_paragraph(doc, f"الهاتف: {patient.phone}")
    _add_rtl_paragraph(doc, f"العنوان: {patient.permanent_address}")
    if patient.emergency_contact_name:
        _add_rtl_paragraph(doc, f"جهة اتصال للطوارئ: {patient.emergency_contact_name} ({patient.emergency_contact_relation}) - {patient.emergency_contact_phone}")
    _add_rtl_paragraph(doc, f"تاريخ التسجيل: {patient.registration_date.strftime('%d/%m/%Y') if patient.registration_date else ''}")

    doc.add_heading('المعلومات المستمرة الحالية', level=2)
    _add_rtl_paragraph(doc, f'الحساسية: {patient.get_allergy_status_display()}')
    _add_rtl_paragraph(doc, f'الأدوية المستمرة: {patient.get_medication_status_display()}')
    for row in patient.patientallergy_records.filter(is_active=True):
        _add_rtl_paragraph(doc, f'{row.substance} — {row.reaction} — {row.get_status_display()}')
    for row in patient.patientmedication_records.filter(is_active=True):
        _add_rtl_paragraph(doc, f'{row.name} — {row.dosage} — {row.schedule} — {row.get_status_display()}')
    for row in patient.patientcontact_records.filter(is_active=True):
        _add_rtl_paragraph(doc, f'جهة اتصال: {row.name} — {row.relationship} — {row.phone}')
    # Medical history
    doc.add_heading('المعلومات الطبية الهامة', level=2)
    _add_rtl_paragraph(doc, f"التاريخ العائلي: {patient.family_history or 'غير مسجل'}")
    _add_rtl_paragraph(doc, f"ملاحظات هامة: {patient.important_notes or 'غير مسجل'}")

    # Visits
    doc.add_heading('الزيارات الطبية', level=2)
    visits = patient.visits.filter(deleted_at__isnull=True).prefetch_related(
        'visit_diagnoses__diagnosis', 'visit_medications__medication', 'scale_responses'
    ).order_by('-visit_date')
    if visits.exists():
        for visit in visits:
            doc.add_heading(f"تاريخ الزيارة: {visit.visit_date} (العمر: {visit.age_at_visit() if hasattr(visit, 'age_at_visit') else ''})", level=3)
            _add_rtl_paragraph(doc, f"الشكوى الرئيسية: {visit.main_complaints}")
            _add_rtl_paragraph(doc, f"القصة المرضية: {visit.history_presenting_complaint}")

            diagnoses = visit.get_diagnoses()
            if diagnoses:
                diag_text = "\n".join(f"- {d.get('arabic_name') or d.get('name', '')}" for d in diagnoses)
                _add_rtl_paragraph(doc, "التشخيصات:\n" + diag_text)

            medications = visit.get_medications()
            if medications:
                med_text = "\n".join(f"- {m.get('name', '')} {m.get('dosage', '')}" for m in medications)
                _add_rtl_paragraph(doc, "الأدوية:\n" + med_text)

            lab_values = visit.get_lab_values()
            if lab_values:
                lab_text = "\n".join(
                    f"- {lv.get('name', '')}: {lv.get('value', '')} {lv.get('unit', '')} "
                    f"({lv.get('reference_range', '—')}; {lv.get('status', '—')}; {lv.get('date', '—')})"
                    for lv in lab_values
                )
                _add_rtl_paragraph(doc, "التحاليل:\n" + lab_text)

            for scale_response in visit.scale_responses.all():
                response_values = scale_response.responses_json or {}
                _add_rtl_paragraph(
                    doc,
                    f"المقياس {scale_response.scale_name_snapshot}: "
                    f"الدرجة {response_values.get('__score', 'غير محدد')}",
                )

            _add_rtl_paragraph(doc, f"حالة التوثيق: {visit.status}")
            _add_rtl_paragraph(doc, f"الحالة السريرية: {visit.clinical_status or 'غير محدد'}")
            _add_rtl_paragraph(doc, f"مرافقة: {visit.accompanied_by} {visit.companion_relation or ''}")
            _add_rtl_paragraph(doc, f"التوصيات العلاجية: {visit.treatment_text}")
            _add_rtl_paragraph(doc, f"ملاحظات الطبيب: {visit.doctor_notes}")
    else:
        _add_rtl_paragraph(doc, "لا توجد زيارات مسجلة.")

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


def _add_rtl_paragraph(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run(text)
    run.font.name = 'Arial'
    return p
