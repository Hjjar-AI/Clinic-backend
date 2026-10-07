import csv
import io
from django.utils import timezone
from datetime import datetime
from django.db import transaction
from django.core.exceptions import ValidationError
from apps.patients.models import Patient
from apps.clinical.models import DiagnosisOption, MedicationOption
from apps.accounts.models import User
from core.normalization import normalize_digits, normalize_identifier, normalize_name, normalize_phone

class TableRow(dict):
    @property
    def iloc(self):
        return list(self.values())

class ImportTable:
    def __init__(self, columns, rows):
        self.columns, self.rows = columns, rows
    def iterrows(self):
        return enumerate(TableRow(zip(self.columns, row)) for row in self.rows)

class BulkImportService:
    def _read_file(self, file_stream):
        import zipfile
        try:
            return self._parse_file(file_stream)
        except (zipfile.BadZipFile, ValueError, UnicodeError) as exc:
            raise ValidationError('تعذر قراءة الملف؛ يلزم CSV أو XLSX صالح') from exc

    def _parse_file(self, file_stream):
        file_stream.seek(0)
        raw = file_stream.read()
        if raw.startswith(b'PK'):
            import zipfile
            from django.conf import settings
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(info.file_size for info in archive.infolist()) > getattr(settings, 'MAX_BULK_IMPORT_SIZE', 200*1024*1024)*5:
                    raise ValidationError('الملف بعد فك الضغط يتجاوز الحد المسموح')
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                values = wb.active.iter_rows(values_only=True)
                headers = next(values, None)
                rows = []
                if headers is None:
                    raise ValidationError('الملف فارغ')
                for row in values:
                    if len(rows) >= 100_000:
                        raise ValidationError('عدد الصفوف يتجاوز الحد المسموح')
                    if any(isinstance(v, str) and v.startswith('=') for v in row):
                        raise ValidationError('استبدل الصيغ في ملف الاستيراد بقيم نصية')
                    rows.append(['' if v is None else v.isoformat() if isinstance(v, datetime) else str(v) for v in row])
            finally:
                wb.close()
        else:
            try:
                rows = list(csv.reader(io.StringIO(raw.decode('utf-8-sig'))))
            except UnicodeError as exc:
                raise ValidationError('يلزم ملف CSV بترميز UTF-8 أو XLSX') from exc
            headers = rows.pop(0) if rows else None
        if not headers or not rows or len(rows) > 100_000:
            raise ValidationError('الملف فارغ أو يتجاوز الحد المسموح')
        columns = [str(h or '').strip() for h in headers]
        if len(set(columns)) != len(columns) or not all(columns):
            raise ValidationError('عناوين الأعمدة مفقودة أو مكررة')
        if any(len(row) != len(columns) for row in rows):
            raise ValidationError('عدد الأعمدة غير متسق')
        return ImportTable(columns, rows)

    def _patient_kwargs(self, row, doctor_id, created_by_id):
        national_id = normalize_identifier(row.get('الرقم الوطني', ''))
        raw_date = (normalize_digits(row.get('تاريخ الإضافة', '')) or '').strip()
        admission_date = None
        if raw_date:
            for date_format in ('%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%d/%m/%Y', '%d-%m-%Y'):
                try:
                    admission_date = datetime.strptime(raw_date, date_format).date()
                    break
                except ValueError:
                    continue
            if admission_date is None:
                raise ValueError('تاريخ الإضافة غير صالح')
        raw_dob_year = (normalize_digits(row.get('سنة الميلاد', '')) or '').strip()
        dob_year = int(raw_dob_year) if raw_dob_year else None
        if dob_year is not None and not (1900 <= dob_year <= timezone.localdate().year):
            raise ValueError('سنة الميلاد غير صالحة')
        return {
            'first_name': normalize_name(row.get('الاسم الأول', '')),
            'father_name': normalize_name(row.get('اسم الأب', '')) or None,
            'surname': normalize_name(row.get('اللقب', '')) or None,
            'mother_name': normalize_name(row.get('اسم الأم', '')) or None,
            'dob_year': dob_year,
            'gender': row.get('الجنس', ''),
            'national_id': national_id,
            'marital_status': row.get('الحالة الاجتماعية', ''),
            'occupation': row.get('المهنة', ''),
            'phone': normalize_phone(row.get('الهاتف', '')),
            'permanent_address': row.get('العنوان', ''),
            'emergency_contact_name': row.get('جهة اتصال للطوارئ (الاسم)', ''),
            'emergency_contact_relation': row.get('صلة القرابة', ''),
            'emergency_contact_phone': normalize_phone(row.get('هاتف جهة الاتصال', '')),
            'family_history': row.get('التاريخ العائلي', ''),
            'important_notes': row.get('ملاحظات هامة', ''),
            'admission_date': admission_date,
            'doctor_id': doctor_id,
            'created_by_id': created_by_id,
        }

    def _patients(self, file_stream, doctor_id, created_by_id, execute=False):
        if not User.objects.filter(pk=doctor_id, is_active=True, role__in=['doctor','admin']).exists():
            raise ValidationError('الطبيب غير صالح')
        results = {'added': 0, 'skipped': 0, 'failed': 0, 'errors': [], 'warnings': [], 'rows': []}
        seen_ids, seen_phones = set(), set()
        for idx, source in self._read_file(file_stream).iterrows():
            row = {k.strip(): v.strip() for k, v in source.items()}
            national_id = normalize_identifier(row.get('الرقم الوطني', ''))
            phone = normalize_phone(row.get('الهاتف', ''))
            if national_id and (national_id in seen_ids or Patient.objects.filter(national_id=national_id).exists()):
                results['skipped'] += 1
                results['rows'].append({'row': idx+2, 'outcome': 'duplicate', 'reason': 'الرقم الوطني مكرر'})
                continue
            if phone and (phone in seen_phones or Patient.objects.filter(phone=phone).exists()):
                results['warnings'].append(f'الصف {idx+2}: الهاتف مشترك مع مريض آخر؛ لن يتم حذف الصف')
            try:
                kwargs = self._patient_kwargs(row, doctor_id, created_by_id)
                from apps.patients.services import PatientService
                PatientService()._validate(kwargs)
                patient = Patient(**kwargs)
                patient.full_clean()
                if execute:
                    with transaction.atomic():
                        patient.save()
                results['added'] += 1
                if national_id: seen_ids.add(national_id)
                if phone: seen_phones.add(phone)
                results['rows'].append({'row': idx+2, 'outcome': 'created' if execute else 'ready', 'reason': ''})
            except (ValidationError, ValueError, TypeError) as exc:
                results['failed'] += 1
                results['errors'].append(f'الصف {idx+2}: {exc}')
                results['rows'].append({'row': idx+2, 'outcome': 'failed', 'reason': str(exc)})
        return results

    def preview_patients(self, file_stream, doctor_id, created_by_id):
        return self._patients(file_stream, doctor_id, created_by_id)

    @transaction.atomic
    def import_patients(self, file_stream, doctor_id, created_by_id):
        return self._patients(file_stream, doctor_id, created_by_id, True)

class OptionsImportService(BulkImportService):
    def _catalog_rows(self, file_stream, kind, column_map=None):
        model = DiagnosisOption if kind == 'diagnosis' else MedicationOption
        keys = ('code', 'english_name', 'arabic_name') if kind == 'diagnosis' else ('generic_english', 'generic_arabic', 'dosage', 'brand_english', 'brand_arabic')
        table = self._read_file(file_stream)
        if column_map is not None:
            if not isinstance(column_map, dict) or any(key not in keys for key in column_map):
                raise ValidationError('خريطة أعمدة غير صالحة')
            for value in column_map.values():
                if value not in (None, '') and (type(value) is not int or not 0 <= value < len(table.columns)):
                    raise ValidationError('رقم العمود خارج نطاق الملف')
        result, seen = [], set()
        for idx, row in table.iterrows():
            try:
                data = {}
                for index, key in enumerate(keys):
                    mapped = column_map.get(key) if column_map is not None else index
                    data[key] = row.iloc[mapped].strip() if type(mapped) is int and mapped < len(row) else ''
                required = keys if kind == 'diagnosis' else ('generic_english',)
                if any(not data[key] for key in required):
                    raise ValidationError('الأعمدة المطلوبة مفقودة')
                identity = (data['code'],) if kind == 'diagnosis' else tuple(data[k].casefold() for k in ('generic_english','dosage','brand_english'))
                if identity in seen:
                    result.append({'row': idx+2, 'outcome': 'duplicate', 'reason': 'صف مكرر'})
                    continue
                lookup = {'code': data['code']} if kind == 'diagnosis' else {k+'__iexact': data[k] for k in ('generic_english','dosage','brand_english')}
                existing = model.all_objects.filter(**lookup).first()
                was_inactive = existing and (existing.deleted_at is not None or not existing.is_active)
                instance = existing or model()
                for key, value in data.items(): setattr(instance,key,value)
                instance.is_active=True
                instance.deleted_at=None
                instance.full_clean()
                seen.add(identity)
                result.append({'row':idx+2, 'outcome': 'reactivate' if was_inactive else 'update' if existing else 'create',
                    'reason':'', '_data':data, '_pk':existing.pk if existing else None})
            except (ValidationError, ValueError, TypeError) as exc:
                result.append({'row':idx+2, 'outcome':'failed', 'reason':str(exc)})
        return result

    @staticmethod
    def _summarize(rows):
        counts = {key:sum(r['outcome']==key for r in rows) for key in ('create','update','reactivate','duplicate','failed')}
        return {'rows':[{k:v for k,v in r.items() if not k.startswith('_')} for r in rows], 'counts':counts, 'total':len(rows)}

    def preview_diagnoses(self, file_stream, merge_mode='merge'):
        rows = self._catalog_rows(file_stream, 'diagnosis')
        result = self._summarize(rows)
        valid_ids = [r['_pk'] for r in rows if r.get('_pk')]
        result['will_retire_existing'] = DiagnosisOption.objects.exclude(pk__in=valid_ids).count() if merge_mode == 'replace' else 0
        return result

    def preview_medications(self, file_stream, column_map, merge_mode='merge'):
        rows = self._catalog_rows(file_stream, 'medication', column_map)
        result = self._summarize(rows)
        valid_ids = [r['_pk'] for r in rows if r.get('_pk')]
        result['will_retire_existing'] = MedicationOption.objects.exclude(pk__in=valid_ids).count() if merge_mode == 'overwrite' else 0
        return result

    @transaction.atomic
    def _import_catalog(self, file_stream, kind, column_map=None, merge_mode='merge'):
        allowed = {'merge','replace'} if kind == 'diagnosis' else {'merge','overwrite'}
        if merge_mode not in allowed: raise ValidationError('طريقة الدمج غير صالحة')
        rows = self._catalog_rows(file_stream, kind, column_map)
        valid = [r for r in rows if '_data' in r]
        if not valid: raise ValidationError('لا توجد صفوف صالحة؛ لم تتغير القائمة')
        # Replacement is all-or-nothing; never retire records for a partially invalid file.
        if merge_mode != 'merge' and any(r['outcome']=='failed' for r in rows):
            raise ValidationError('صحح جميع الصفوف قبل استبدال القائمة')
        model = DiagnosisOption if kind == 'diagnosis' else MedicationOption
        list(model.all_objects.select_for_update().order_by('pk'))
        if merge_mode != 'merge':
            for instance in model.objects.all(): instance.soft_delete(reason='Catalog replacement import')
        result = {'added':0,'updated':0,'reactivated':0,'rows':[]}
        for row in rows:
            if '_data' not in row:
                result['rows'].append({k:v for k,v in row.items() if not k.startswith('_')})
                continue
            with transaction.atomic():
                instance = model.all_objects.get(pk=row['_pk']) if row['_pk'] else model()
                was_inactive = not instance.is_active or instance.deleted_at is not None
                for key,value in row['_data'].items(): setattr(instance,key,value)
                instance.is_active=True
                instance.deleted_at=None
                instance.version += int(instance.pk is not None)
                instance.full_clean()
                instance.save()
                outcome = 'reactivated' if row['outcome'] == 'reactivate' else 'updated' if row['_pk'] else 'added'
                result[outcome] += 1
                result['rows'].append({'row':row['row'], 'outcome':outcome, 'reason':''})
        return result

    def import_diagnoses(self, file_stream, merge_mode='merge'):
        return self._import_catalog(file_stream, 'diagnosis', merge_mode=merge_mode)

    def import_medications(self, file_stream, column_map=None, merge_mode='merge'):
        return self._import_catalog(file_stream, 'medication', column_map, merge_mode)
